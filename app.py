"""
DNS Panel - lightweight query-log viewer for mosdns.

Tails the mosdns query_summary log (JSON-suffixed lines), stores parsed
records in SQLite, and exposes a small REST API + dashboard so you can look
up, per client IP, which domains were resolved. Includes a simple
username/password login, an IP -> contact-name address book, and CSV /
TXT / XLSX / PDF export.

Env vars:
  LOG_PATH         path to mosdns.log            (default /var/log/mosdns/mosdns.log)
  DB_PATH          path to sqlite db file         (default /data/dns_panel.db)
  RETENTION_DAYS   delete records older than this (default 30, 0 = keep forever)
  PORT             http port                      (default 8080)
  PANEL_USERNAME   login username                 (default admin)
  PANEL_PASSWORD   login password                 (default changeme -- CHANGE THIS)
  SESSION_SECRET   fixed session signing key       (default: auto-generated & persisted in the db)

  VPN sync (optional -- auto-fill the contacts book from a WireGuard admin panel):
  VPN_PANEL_TYPE            "wgdashboard" | "wg-easy" | unset (default: unset = disabled)
  VPN_PANEL_URL             base URL of the panel, e.g. http://10.0.0.1:10086
  VPN_PANEL_API_KEY         (wgdashboard only) API key from Settings -> API Key
  VPN_PANEL_INTERFACE       (wgdashboard only) WireGuard interface/config name, e.g. wg0
  VPN_PANEL_USERNAME        (wg-easy only) admin username
  VPN_PANEL_PASSWORD        (wg-easy only) admin password
  VPN_SYNC_INTERVAL_SECONDS how often to poll                (default 300 = 5 minutes)
"""

import csv
import gzip
import io
import ipaddress
import json
import os
import re
import secrets
import shutil
import sqlite3
import threading
import time

import jdatetime
import requests
from contextlib import asynccontextmanager
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware

# mosdns logs its timestamps in this timezone (Iran does not observe DST,
# so this is a fixed +03:30 offset). All comparisons, exports and generated
# dates in the panel are done in this timezone so they line up with the log.
LOG_TZ = ZoneInfo(os.environ.get("LOG_TIMEZONE", "Asia/Tehran"))


def tehran_now():
    return datetime.now(LOG_TZ)


def tehran_ts_str(dt):
    """Format a datetime as a bare 'YYYY-MM-DDTHH:MM:SS.mmm' string (no
    offset) -- this is directly comparable with the stored `ts` column
    because it shares the same prefix as the offset-suffixed log timestamps
    (e.g. '...15:05:31.314+0330'); a shorter matching prefix always sorts
    before/equal to the longer string in SQLite's lexicographic ordering."""
    return dt.strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3]


LOG_PATH = os.environ.get("LOG_PATH", "/var/log/mosdns/mosdns.log")
DB_PATH = os.environ.get("DB_PATH", "/data/dns_panel.db")
RETENTION_DAYS = int(os.environ.get("RETENTION_DAYS", "30"))

# Once the live mosdns.log grows past this size, its current contents are
# gzip-compressed into LOG_ARCHIVE_DIR and the live file is truncated in
# place (never renamed/deleted -- mosdns keeps writing to the same file
# descriptor). Nothing is ever lost; only you clean up the archive folder.
LOG_ARCHIVE_DIR = os.environ.get("LOG_ARCHIVE_DIR", "/log-archive")
LOG_MAX_SIZE_MB = int(os.environ.get("LOG_MAX_SIZE_MB", "200"))
PANEL_USERNAME = os.environ.get("PANEL_USERNAME", "admin")
PANEL_PASSWORD = os.environ.get("PANEL_PASSWORD", "changeme")

VPN_PANEL_TYPE = (os.environ.get("VPN_PANEL_TYPE") or "").strip().lower()
VPN_PANEL_URL = (os.environ.get("VPN_PANEL_URL") or "").rstrip("/")
VPN_PANEL_API_KEY = os.environ.get("VPN_PANEL_API_KEY", "")
VPN_PANEL_INTERFACE = os.environ.get("VPN_PANEL_INTERFACE", "")
VPN_PANEL_USERNAME = os.environ.get("VPN_PANEL_USERNAME", "")
VPN_PANEL_PASSWORD = os.environ.get("VPN_PANEL_PASSWORD", "")
VPN_SYNC_INTERVAL_SECONDS = int(os.environ.get("VPN_SYNC_INTERVAL_SECONDS", "300"))

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
FONT_DIR = os.path.join(BASE_DIR, "fonts")

# mosdns (zap logger) query_summary line looks like:
# 2026-09-10T00:58:53.773+0330 INFO main_sequence.r0 main {"uqid": 1, "client": "::ffff:172.20.20.33",
#   "qname": "raw.githubusercontent.com.", "qtype": 1, "qclass": 1, "rcode": 0, "elapsed": "518ms"}
#
# The 4th field is whatever title you gave query_summary in config.yaml, so we
# don't hard-code it -- grab everything up to the trailing JSON blob and
# validate the JSON itself has the fields we need.
LINE_RE = re.compile(
    r"^(?P<ts>\S+)\s+(?P<level>\S+)\s+(?P<logger>\S+)\s+(?P<msg>.*?)\s+(?P<json>\{.*\})\s*$"
)
REQUIRED_JSON_KEYS = {"client", "qname"}

QTYPE_MAP = {
    1: "A", 2: "NS", 5: "CNAME", 6: "SOA", 12: "PTR", 15: "MX", 16: "TXT",
    28: "AAAA", 33: "SRV", 41: "OPT", 43: "DS", 46: "RRSIG", 47: "NSEC",
    48: "DNSKEY", 52: "TLSA", 64: "SVCB", 65: "HTTPS", 255: "ANY",
}
RCODE_MAP = {
    0: "NOERROR", 1: "FORMERR", 2: "SERVFAIL", 3: "NXDOMAIN",
    4: "NOTIMP", 5: "REFUSED",
}

ELAPSED_RE = re.compile(r"([\d.]+)(ns|µs|us|ms|s)")


def parse_elapsed_ms(s):
    if not s:
        return None
    m = ELAPSED_RE.match(s.strip())
    if not m:
        return None
    val, unit = float(m.group(1)), m.group(2)
    factor = {"ns": 1e-6, "us": 1e-3, "µs": 1e-3, "ms": 1, "s": 1000}.get(unit, 1)
    return round(val * factor, 3)


def normalize_ip(client):
    if not client:
        return client
    client = client.strip()
    if client.startswith("::ffff:"):
        client = client[len("::ffff:"):]
    return client


def parse_line(line):
    m = LINE_RE.match(line.strip())
    if not m:
        return None
    try:
        data = json.loads(m.group("json"))
    except json.JSONDecodeError:
        return None
    if not REQUIRED_JSON_KEYS.issubset(data.keys()):
        return None  # some other log line that happens to end in { ... }
    qname = data.get("qname", "").rstrip(".")
    if not qname:
        return None
    return {
        "ts": m.group("ts"),
        "client_ip": normalize_ip(data.get("client")),
        "qname": qname,
        "qtype": data.get("qtype"),
        "rcode": data.get("rcode"),
        "elapsed_ms": parse_elapsed_ms(data.get("elapsed")),
    }


def ip_class(ip):
    """172.x = our own docker/server-side addresses, 10.x = real LAN users."""
    if ip.startswith("172."):
        return "server"
    if ip.startswith("10."):
        return "user"
    return "other"


# ---------------------------------------------------------------------------
# Storage
# ---------------------------------------------------------------------------

def get_conn():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.row_factory = sqlite3.Row
    return conn


def ensure_schema():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = get_conn()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS queries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ts TEXT NOT NULL,
            client_ip TEXT NOT NULL,
            qname TEXT NOT NULL,
            qtype INTEGER,
            rcode INTEGER,
            elapsed_ms REAL
        );
        CREATE INDEX IF NOT EXISTS idx_q_client ON queries(client_ip);
        CREATE INDEX IF NOT EXISTS idx_q_qname ON queries(qname);
        CREATE INDEX IF NOT EXISTS idx_q_ts ON queries(ts);
        CREATE INDEX IF NOT EXISTS idx_q_client_qname_ts ON queries(client_ip, qname, ts);

        CREATE TABLE IF NOT EXISTS ingest_state (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            file_path TEXT,
            file_offset INTEGER,
            file_inode INTEGER
        );

        CREATE TABLE IF NOT EXISTS contacts (
            ip TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            source TEXT NOT NULL DEFAULT 'manual'
        );

        CREATE TABLE IF NOT EXISTS app_meta (
            key TEXT PRIMARY KEY,
            value TEXT
        );
        """
    )
    # migration for databases created before the `source` column existed
    cols = [r["name"] for r in conn.execute("PRAGMA table_info(contacts)").fetchall()]
    if "source" not in cols:
        conn.execute("ALTER TABLE contacts ADD COLUMN source TEXT NOT NULL DEFAULT 'manual'")
    conn.commit()
    conn.close()


def get_or_create_session_secret():
    env_secret = os.environ.get("SESSION_SECRET")
    if env_secret:
        return env_secret
    conn = get_conn()
    row = conn.execute("SELECT value FROM app_meta WHERE key='session_secret'").fetchone()
    if row:
        val = row["value"]
    else:
        val = secrets.token_hex(32)
        conn.execute("INSERT INTO app_meta (key, value) VALUES ('session_secret', ?)", (val,))
        conn.commit()
    conn.close()
    return val


def load_state():
    conn = get_conn()
    row = conn.execute("SELECT * FROM ingest_state WHERE id=1").fetchone()
    conn.close()
    if row is None:
        return {"file_path": None, "file_offset": 0, "file_inode": None}
    return dict(row)


def save_state(path, offset, inode):
    conn = get_conn()
    conn.execute(
        """INSERT INTO ingest_state (id, file_path, file_offset, file_inode)
           VALUES (1, ?, ?, ?)
           ON CONFLICT(id) DO UPDATE SET file_path=excluded.file_path,
               file_offset=excluded.file_offset, file_inode=excluded.file_inode""",
        (path, offset, inode),
    )
    conn.commit()
    conn.close()


def insert_batch(records):
    if not records:
        return
    conn = get_conn()
    conn.executemany(
        """INSERT INTO queries (ts, client_ip, qname, qtype, rcode, elapsed_ms)
           VALUES (:ts, :client_ip, :qname, :qtype, :rcode, :elapsed_ms)""",
        records,
    )
    conn.commit()
    conn.close()


# ---------------------------------------------------------------------------
# Background log tailer
# ---------------------------------------------------------------------------

def tail_loop():
    while True:
        try:
            _tail_once()
            _archive_if_oversized()
        except Exception as e:  # keep the thread alive no matter what
            print(f"[tailer] error: {e}", flush=True)
        time.sleep(1)


def _tail_once():
    if not os.path.exists(LOG_PATH):
        return
    st = os.stat(LOG_PATH)
    state = load_state()
    offset = state["file_offset"] or 0
    inode = state["file_inode"]

    if inode is not None and (inode != st.st_ino or st.st_size < offset):
        offset = 0  # rotated or truncated

    if st.st_size == offset and inode == st.st_ino:
        return

    batch = []
    with open(LOG_PATH, "r", encoding="utf-8", errors="replace") as f:
        f.seek(offset)
        for line in f:
            rec = parse_line(line)
            if rec:
                batch.append(rec)
        new_offset = f.tell()

    insert_batch(batch)
    save_state(LOG_PATH, new_offset, st.st_ino)


def _archive_if_oversized():
    """If the live log has grown past LOG_MAX_SIZE_MB, gzip its current
    contents into LOG_ARCHIVE_DIR and truncate it *in place* (same inode,
    same file descriptor mosdns is already writing to -- never renamed or
    recreated, so mosdns keeps working without a restart). Only called
    right after _tail_once(), so everything currently in the file has
    already been imported into the database before it's archived."""
    if LOG_MAX_SIZE_MB <= 0 or not os.path.exists(LOG_PATH):
        return
    if os.path.getsize(LOG_PATH) < LOG_MAX_SIZE_MB * 1024 * 1024:
        return

    os.makedirs(LOG_ARCHIVE_DIR, exist_ok=True)
    stamp = tehran_now().strftime("%Y%m%d-%H%M%S")
    archive_path = os.path.join(LOG_ARCHIVE_DIR, f"mosdns-{stamp}.log.gz")

    with open(LOG_PATH, "rb") as src, gzip.open(archive_path, "wb") as dst:
        shutil.copyfileobj(src, dst)

    # truncate in place -- do NOT delete/rename the file
    with open(LOG_PATH, "r+", encoding="utf-8") as f:
        f.truncate(0)

    inode = os.stat(LOG_PATH).st_ino
    save_state(LOG_PATH, 0, inode)
    print(f"[archive] compressed log to {archive_path}, truncated live file", flush=True)


def rebuild_from_archive_if_empty():
    """One-time startup step: if the database has no queries at all (e.g.
    you deleted data/ to start fresh) but LOG_ARCHIVE_DIR has old archived
    logs, replay them all back into the database before live tailing
    starts, so nothing you archived is missing from the panel."""
    try:
        conn = get_conn()
        count = conn.execute("SELECT COUNT(*) c FROM queries").fetchone()["c"]
        conn.close()
        if count > 0:
            return
        if not os.path.isdir(LOG_ARCHIVE_DIR):
            return
        archives = sorted(
            f for f in os.listdir(LOG_ARCHIVE_DIR) if f.endswith(".log.gz")
        )
        if not archives:
            return

        print(f"[archive] database is empty, replaying {len(archives)} archived log file(s)...", flush=True)
        total = 0
        for name in archives:
            path = os.path.join(LOG_ARCHIVE_DIR, name)
            batch = []
            with gzip.open(path, "rt", encoding="utf-8", errors="replace") as f:
                for line in f:
                    rec = parse_line(line)
                    if rec:
                        batch.append(rec)
            insert_batch(batch)
            total += len(batch)
            print(f"[archive]   {name}: {len(batch)} records", flush=True)
        print(f"[archive] replay done, {total} records restored from archive", flush=True)
    except Exception as e:
        print(f"[archive] rebuild-from-archive failed: {e}", flush=True)


def retention_loop():
    if RETENTION_DAYS <= 0:
        return
    while True:
        try:
            cutoff = tehran_ts_str(tehran_now() - timedelta(days=RETENTION_DAYS))
            conn = get_conn()
            conn.execute("DELETE FROM queries WHERE ts < ?", (cutoff,))
            conn.commit()
            conn.execute("VACUUM")
            conn.close()
        except Exception as e:
            print(f"[retention] error: {e}", flush=True)
        time.sleep(6 * 3600)  # every 6h


# ---------------------------------------------------------------------------
# VPN panel sync -- auto-fill the contacts book from a WireGuard admin panel.
#
# Design: each panel type is a small adapter function that returns either
#   - a list of {"name": ..., "ip": ...} dicts (success), or
#   - None (disabled / misconfigured / any error) -- the caller must then do
#     nothing at all, leaving the contacts book exactly as it is.
#
# To support a new panel later (e.g. a personal WireGuard panel), add another
# adapter function with this same contract and one more branch in
# fetch_vpn_peers(); nothing else in the sync engine needs to change.
# ---------------------------------------------------------------------------

def _first_ip(raw):
    """Peer 'allowed ip' / 'address' fields come back as a single CIDR
    string, a comma-separated list of CIDRs (IPv4+IPv6 mixed), or sometimes
    a list. Return the first valid bare IPv4 address (no /mask), skipping
    any IPv6 entries -- mirrors the confirmed-working extraction logic."""
    if not raw:
        return ""
    items = raw if isinstance(raw, list) else str(raw).split(",")
    for item in items:
        item = str(item).strip()
        if not item:
            continue
        try:
            network = ipaddress.ip_network(item, strict=False)
        except ValueError:
            continue
        if network.version == 4:
            return str(network.network_address)
    return ""


def fetch_peers_wgdashboard():
    if not VPN_PANEL_URL or not VPN_PANEL_API_KEY or not VPN_PANEL_INTERFACE:
        return None
    url = f"{VPN_PANEL_URL}/api/getWireguardConfigurationInfo"
    try:
        resp = requests.get(
            url,
            params={"configurationName": VPN_PANEL_INTERFACE},
            headers={"wg-dashboard-apikey": VPN_PANEL_API_KEY},
            timeout=8,
        )
        resp.raise_for_status()
        payload = resp.json()
    except Exception as e:
        print(f"[vpn_sync:wgdashboard] request failed: {e}", flush=True)
        return None

    if not payload.get("status"):
        print(f"[vpn_sync:wgdashboard] API returned status=false: {payload.get('message')}", flush=True)
        return None

    data = payload.get("data")
    # be tolerant of the exact shape (varies a bit across WGDashboard versions)
    peer_list = None
    if isinstance(data, list):
        peer_list = data
    elif isinstance(data, dict):
        peer_list = data.get("configurationPeers") or data.get("peers") or data.get("Peers")
    if peer_list is None:
        print("[vpn_sync:wgdashboard] unexpected response shape, no peer list found", flush=True)
        return None

    peers = []
    for p in peer_list:
        name = (p.get("name") or p.get("Name") or "").strip()
        ip = _first_ip(p.get("allowed_ip") or p.get("allowed_ips") or p.get("AllowedIPs"))
        if name and ip:
            peers.append({"name": name, "ip": ip})
    return peers


def fetch_peers_wgeasy():
    if not VPN_PANEL_URL or not VPN_PANEL_USERNAME or not VPN_PANEL_PASSWORD:
        return None
    url = f"{VPN_PANEL_URL}/api/client"
    try:
        resp = requests.get(
            url, auth=(VPN_PANEL_USERNAME, VPN_PANEL_PASSWORD), timeout=8
        )
        if resp.status_code == 401:
            print("[vpn_sync:wg-easy] login failed (401) -- check username/password", flush=True)
            return None
        resp.raise_for_status()
        data = resp.json()
    except Exception as e:
        print(f"[vpn_sync:wg-easy] request failed: {e}", flush=True)
        return None

    client_list = data if isinstance(data, list) else (data.get("clients") or data.get("data"))
    if client_list is None:
        print("[vpn_sync:wg-easy] unexpected response shape, no client list found", flush=True)
        return None

    peers = []
    for c in client_list:
        name = (c.get("name") or "").strip()
        # wg-easy exposes the assigned tunnel IP as `ipv4Address`; keep the
        # older/generic field names as a fallback for other wg-easy versions.
        ip = _first_ip(c.get("ipv4Address") or c.get("address") or c.get("Address"))
        if name and ip:
            peers.append({"name": name, "ip": ip})
    return peers


def fetch_vpn_peers():
    if VPN_PANEL_TYPE == "wgdashboard":
        return fetch_peers_wgdashboard()
    if VPN_PANEL_TYPE == "wg-easy":
        return fetch_peers_wgeasy()
    return None  # unset / unknown panel type -> disabled


def sync_vpn_contacts():
    peers = fetch_vpn_peers()
    if peers is None:
        return  # disabled or unreachable/misconfigured -- touch nothing

    conn = get_conn()
    existing = {r["ip"]: r for r in conn.execute("SELECT ip, name, source FROM contacts").fetchall()}
    seen_ips = set()
    for p in peers:
        ip, name = p["ip"], p["name"]
        seen_ips.add(ip)
        cur = existing.get(ip)
        if cur is None:
            conn.execute(
                "INSERT INTO contacts (ip, name, source) VALUES (?, ?, 'synced')", (ip, name)
            )
        elif cur["source"] == "synced" and cur["name"] != name:
            conn.execute("UPDATE contacts SET name = ? WHERE ip = ?", (name, ip))
        # source == 'manual' -> never touched, manual entries always win

    for ip, row in existing.items():
        if row["source"] == "synced" and ip not in seen_ips:
            conn.execute("DELETE FROM contacts WHERE ip = ?", (ip,))

    conn.commit()
    conn.close()


def vpn_sync_loop():
    if not VPN_PANEL_TYPE:
        return  # not configured at all -- don't even start polling
    while True:
        try:
            sync_vpn_contacts()
        except Exception as e:
            print(f"[vpn_sync] error: {e}", flush=True)
        time.sleep(VPN_SYNC_INTERVAL_SECONDS)


# ---------------------------------------------------------------------------
# Contacts helpers
# ---------------------------------------------------------------------------

def get_contacts_map(conn=None):
    close = False
    if conn is None:
        conn = get_conn()
        close = True
    m = {r["ip"]: r["name"] for r in conn.execute("SELECT ip, name FROM contacts").fetchall()}
    if close:
        conn.close()
    return m


def enrich_client(row, contacts):
    d = dict(row)
    d["display_name"] = contacts.get(d["client_ip"], d["client_ip"])
    d["ip_class"] = ip_class(d["client_ip"])
    return d


def enrich_query_row(row, contacts):
    d = dict(row)
    d["qtype_name"] = QTYPE_MAP.get(d.get("qtype"), str(d.get("qtype")))
    d["rcode_name"] = RCODE_MAP.get(d.get("rcode"), str(d.get("rcode")))
    d["display_name"] = contacts.get(d["client_ip"], d["client_ip"])
    d["ip_class"] = ip_class(d["client_ip"])
    return d


# ---------------------------------------------------------------------------
# Query helpers (shared by JSON API + export)
# ---------------------------------------------------------------------------

def normalize_date_bound(value, is_end):
    """Turn a <input type=datetime-local> value ('YYYY-MM-DDTHH:MM') into a
    string comparable with our stored ts values ('YYYY-MM-DDTHH:MM:SS.mmm+HHMM').
    Since both share the same prefix format, plain string comparison works."""
    if not value:
        return None
    v = value.strip()
    if len(v) == 16:  # YYYY-MM-DDTHH:MM
        v += ":59.999" if is_end else ":00.000"
    return v


def date_where(date_from, date_to, params, alias=""):
    clauses = []
    lo = normalize_date_bound(date_from, False)
    hi = normalize_date_bound(date_to, True)
    col = f"{alias}ts" if alias else "ts"
    if lo:
        clauses.append(f"{col} >= ?")
        params.append(lo)
    if hi:
        clauses.append(f"{col} <= ?")
        params.append(hi)
    return clauses


def query_stats():
    conn = get_conn()
    total = conn.execute("SELECT COUNT(*) c FROM queries").fetchone()["c"]
    clients_n = conn.execute("SELECT COUNT(DISTINCT client_ip) c FROM queries").fetchone()["c"]
    domains_n = conn.execute("SELECT COUNT(DISTINCT qname) c FROM queries").fetchone()["c"]
    since = tehran_ts_str(tehran_now() - timedelta(hours=24))
    last24 = conn.execute("SELECT COUNT(*) c FROM queries WHERE ts >= ?", (since,)).fetchone()["c"]
    top_domains = [
        dict(r) for r in conn.execute(
            "SELECT qname, COUNT(*) count FROM queries GROUP BY qname ORDER BY count DESC LIMIT 15"
        ).fetchall()
    ]
    contacts = get_contacts_map(conn)
    top_clients = [
        enrich_client(r, contacts) for r in conn.execute(
            "SELECT client_ip, COUNT(*) count FROM queries GROUP BY client_ip ORDER BY count DESC LIMIT 15"
        ).fetchall()
    ]
    blocked = conn.execute("SELECT COUNT(*) c FROM queries WHERE rcode IN (3,5)").fetchone()["c"]
    conn.close()
    return {
        "total_queries": total,
        "unique_clients": clients_n,
        "unique_domains": domains_n,
        "last_24h": last24,
        "blocked_or_refused": blocked,
        "top_domains": top_domains,
        "top_clients": top_clients,
    }


def query_client_info(ip):
    """Exact-match lookup for a single client's summary + contact name.
    Deliberately does NOT use the fuzzy 'q' substring search used by
    query_clients(), since e.g. '10.70.70.2' is a substring of
    '10.70.70.29' and substring matching there previously returned the
    wrong client's info on the detail page."""
    conn = get_conn()
    row = conn.execute(
        """SELECT client_ip, COUNT(*) query_count, COUNT(DISTINCT qname) domain_count,
                  MAX(ts) last_seen, MIN(ts) first_seen
           FROM queries WHERE client_ip = ?""",
        (ip,),
    ).fetchone()
    contacts = get_contacts_map(conn)
    conn.close()
    if row and row["query_count"]:
        return enrich_client(row, contacts)
    return {
        "client_ip": ip, "display_name": contacts.get(ip, ip), "ip_class": ip_class(ip),
        "query_count": 0, "domain_count": 0, "first_seen": None, "last_seen": None,
    }


def query_clients(q="", date_from="", date_to="", page=1, page_size=50, paginate=True):
    conn = get_conn()
    params = []
    where = date_where(date_from, date_to, params)
    where_sql = f"WHERE {' AND '.join(where)}" if where else ""
    rows = [
        dict(r) for r in conn.execute(
            f"""SELECT client_ip, COUNT(*) query_count, COUNT(DISTINCT qname) domain_count,
                       MAX(ts) last_seen, MIN(ts) first_seen
                FROM queries {where_sql}
                GROUP BY client_ip ORDER BY last_seen DESC""",
            params,
        ).fetchall()
    ]
    contacts = get_contacts_map(conn)
    conn.close()
    enriched = [enrich_client(r, contacts) for r in rows]
    if q:
        ql = q.lower()
        enriched = [r for r in enriched if ql in r["client_ip"].lower() or ql in r["display_name"].lower()]
    total = len(enriched)
    if not paginate:
        return {"total": total, "page": 1, "page_size": page_size, "rows": enriched[:page_size]}
    start = (page - 1) * page_size
    return {"total": total, "page": page, "page_size": page_size, "rows": enriched[start:start + page_size]}


def query_client_domains(ip, search="", date_from="", date_to="", page=1, page_size=50, paginate=True):
    conn = get_conn()
    params = [ip]
    where = ["client_ip = ?"]
    if search:
        where.append("qname LIKE ?")
        params.append(f"%{search}%")
    where += date_where(date_from, date_to, params)
    where_sql = " AND ".join(where)

    total = conn.execute(
        f"SELECT COUNT(DISTINCT qname) c FROM queries WHERE {where_sql}", params
    ).fetchone()["c"]

    # Correlated subquery to get the rcode of the most recent query for each
    # domain -- scoped to the same client + date range as the outer query,
    # so "last result" always matches the "last_seen" timestamp shown in the
    # same row (not some later query outside the selected date range).
    sub_params = [ip]
    sub_clauses = ["q2.client_ip = ?", "q2.qname = queries.qname"]
    sub_clauses += date_where(date_from, date_to, sub_params, alias="q2.")
    sub_where_sql = " AND ".join(sub_clauses)

    sql = f"""SELECT qname, COUNT(*) count, MAX(ts) last_seen, MIN(ts) first_seen,
                     GROUP_CONCAT(DISTINCT qtype) qtypes,
                     (SELECT q2.rcode FROM queries q2
                      WHERE {sub_where_sql}
                      ORDER BY q2.ts DESC, q2.id DESC LIMIT 1) AS last_rcode
              FROM queries WHERE {where_sql}
              GROUP BY qname ORDER BY last_seen DESC"""
    q_params = sub_params + list(params)
    if paginate:
        sql += " LIMIT ? OFFSET ?"
        q_params += [page_size, (page - 1) * page_size]
    else:
        sql += " LIMIT ?"
        q_params += [page_size]
    rows = [dict(r) for r in conn.execute(sql, q_params).fetchall()]
    conn.close()
    for r in rows:
        r["qtypes"] = [QTYPE_MAP.get(int(t), t) for t in r["qtypes"].split(",")] if r["qtypes"] else []
        r["rcode_name"] = RCODE_MAP.get(r.get("last_rcode"), str(r.get("last_rcode")))
    return {"total": total, "page": page if paginate else 1, "page_size": page_size, "rows": rows}


def query_search(q, date_from="", date_to="", page=1, page_size=50, paginate=True):
    conn = get_conn()
    params = [f"%{q}%"]
    where = ["qname LIKE ?"]
    where += date_where(date_from, date_to, params)
    where_sql = " AND ".join(where)

    total = conn.execute(f"SELECT COUNT(*) c FROM queries WHERE {where_sql}", params).fetchone()["c"]

    sql = f"SELECT ts, client_ip, qname, qtype, rcode, elapsed_ms FROM queries WHERE {where_sql} ORDER BY ts DESC"
    q_params = list(params)
    if paginate:
        sql += " LIMIT ? OFFSET ?"
        q_params += [page_size, (page - 1) * page_size]
    else:
        sql += " LIMIT ?"
        q_params += [page_size]
    rows = conn.execute(sql, q_params).fetchall()
    contacts = get_contacts_map(conn)
    conn.close()
    return {
        "total": total, "page": page if paginate else 1, "page_size": page_size,
        "rows": [enrich_query_row(r, contacts) for r in rows],
    }


def query_recent(date_from="", date_to="", page=1, page_size=50, paginate=True):
    conn = get_conn()
    params = []
    where = date_where(date_from, date_to, params)
    where_sql = f"WHERE {' AND '.join(where)}" if where else ""

    total = conn.execute(f"SELECT COUNT(*) c FROM queries {where_sql}", params).fetchone()["c"]

    sql = f"SELECT ts, client_ip, qname, qtype, rcode, elapsed_ms FROM queries {where_sql} ORDER BY id DESC"
    q_params = list(params)
    if paginate:
        sql += " LIMIT ? OFFSET ?"
        q_params += [page_size, (page - 1) * page_size]
    else:
        sql += " LIMIT ?"
        q_params += [page_size]
    rows = conn.execute(sql, q_params).fetchall()
    contacts = get_contacts_map(conn)
    conn.close()
    return {
        "total": total, "page": page if paginate else 1, "page_size": page_size,
        "rows": [enrich_query_row(r, contacts) for r in rows],
    }


# ---------------------------------------------------------------------------
# Export builders
# ---------------------------------------------------------------------------

COLUMN_LABELS = {
    "client_ip": "آی‌پی", "display_name": "نام/آی‌پی", "ip_class": "نوع",
    "query_count": "تعداد کوئری", "domain_count": "تعداد دامنه",
    "first_seen": "اولین بازدید", "last_seen": "آخرین بازدید",
    "qname": "دامنه", "count": "تعداد", "qtypes": "نوع رکورد",
    "ts": "زمان", "qtype_name": "نوع", "rcode_name": "نتیجه",
    "elapsed_ms": "زمان پاسخ (ms)",
}
IP_CLASS_LABELS_FA = {"server": "سرور", "user": "کاربر", "other": "-"}


def display_ts(ts):
    """Reformat the raw log timestamp ('...+0330') into a clean, human
    readable Tehran-local, Jalali-calendar string for exports (matches what
    the panel itself shows). Falls back to the raw value if parsing fails."""
    if not ts:
        return ts
    try:
        dt = datetime.fromisoformat(ts).astimezone(LOG_TZ)
        jd = jdatetime.datetime.fromgregorian(datetime=dt)
        return jd.strftime("%Y/%m/%d %H:%M:%S")
    except ValueError:
        return ts


def cell_value(row, col):
    v = row.get(col, "")
    if col == "qtypes" and isinstance(v, list):
        return ", ".join(v)
    if col == "ip_class":
        return IP_CLASS_LABELS_FA.get(v, v)
    if col in ("ts", "first_seen", "last_seen"):
        return display_ts(v)
    return "" if v is None else v


def rows_to_csv(rows, columns):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([COLUMN_LABELS.get(c, c) for c in columns])
    for r in rows:
        writer.writerow([cell_value(r, c) for c in columns])
    return "\ufeff" + buf.getvalue()  # BOM so Excel opens UTF-8/Persian correctly


def rows_to_txt(rows, columns):
    lines = ["\t".join(COLUMN_LABELS.get(c, c) for c in columns)]
    for r in rows:
        lines.append("\t".join(str(cell_value(r, c)) for c in columns))
    return "\ufeff" + "\n".join(lines)


def rows_to_xlsx(rows, columns):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append([COLUMN_LABELS.get(c, c) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([cell_value(r, c) for c in columns])
    for col_cells in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells) + 2
        ws.column_dimensions[col_cells[0].column_letter].width = min(width, 40)
    ws.sheet_view.rightToLeft = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_bidi_ready = True
try:
    import arabic_reshaper
    from bidi.algorithm import get_display
except Exception:  # pragma: no cover
    _bidi_ready = False


def fa(text):
    text = str(text)
    if not _bidi_ready:
        return text
    try:
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def jalali_label(value):
    """value is a raw 'YYYY-MM-DDTHH:MM' (Gregorian, Tehran wall-clock) filter
    bound coming from the date-range picker; render it in Jalali for titles."""
    if not value:
        return "..."
    try:
        dt = datetime.fromisoformat(value)
        return jdatetime.datetime.fromgregorian(datetime=dt).strftime("%Y/%m/%d %H:%M")
    except ValueError:
        return value


def rows_to_pdf(rows, columns, title):
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()
    pdf.add_font("Vazir", "", os.path.join(FONT_DIR, "Vazirmatn-Regular.ttf"))
    pdf.add_font("Vazir", "B", os.path.join(FONT_DIR, "Vazirmatn-Bold.ttf"))
    pdf.set_font("Vazir", "B", 16)
    pdf.cell(0, 12, fa(title), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Vazir", "", 9)
    pdf.cell(0, 6, fa(f"تعداد رکورد: {len(rows)} — تاریخ تولید: {jdatetime.datetime.fromgregorian(datetime=tehran_now()).strftime('%Y/%m/%d %H:%M')} (به وقت تهران)"),
              align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    usable_width = pdf.w - pdf.l_margin - pdf.r_margin
    # fpdf2 always lays out table cells left-to-right on the page. Our
    # columns are ordered "reading order" (first = rightmost for a Persian
    # reader), so reverse them here to get the same right-to-left column
    # order the web UI shows, with the first logical column ending up
    # visually on the right.
    columns_rtl = list(reversed(columns))
    widths = [usable_width / max(1, len(columns_rtl))] * len(columns_rtl)

    pdf.set_font("Vazir", "B", 9)
    with pdf.table(col_widths=widths, text_align="CENTER", line_height=6,
                    headings_style=FontFace(family="Vazir", emphasis="B", size_pt=9)) as table:
        header = table.row()
        for c in columns_rtl:
            header.cell(fa(COLUMN_LABELS.get(c, c)))
        pdf.set_font("Vazir", "", 8)
        for r in rows:
            row = table.row()
            for c in columns_rtl:
                val = str(cell_value(r, c))
                if len(val) > 45:
                    val = val[:42] + "..."
                row.cell(fa(val))

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------

def require_auth(request: Request):
    if not request.session.get("auth"):
        raise HTTPException(status_code=401, detail="Unauthorized")


# ---------------------------------------------------------------------------
# App setup
# ---------------------------------------------------------------------------

ensure_schema()
SESSION_SECRET = get_or_create_session_secret()


@asynccontextmanager
async def lifespan(app: FastAPI):
    threading.Thread(target=rebuild_from_archive_if_empty, daemon=True).start()
    threading.Thread(target=tail_loop, daemon=True).start()
    threading.Thread(target=retention_loop, daemon=True).start()
    threading.Thread(target=vpn_sync_loop, daemon=True).start()
    yield


app = FastAPI(title="DNS Panel", lifespan=lifespan)
app.add_middleware(
    SessionMiddleware,
    secret_key=SESSION_SECRET,
    session_cookie="dns_panel_session",
    max_age=14 * 24 * 3600,
)


@app.middleware("http")
async def no_cache_static(request: Request, call_next):
    response = await call_next(request)
    if request.url.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-cache, must-revalidate"
    return response


# ---------------- auth routes ----------------

@app.get("/login")
def login_page():
    return FileResponse(os.path.join(STATIC_DIR, "login.html"))


@app.post("/login")
async def do_login(request: Request):
    form = await request.form()
    username = str(form.get("username", ""))
    password = str(form.get("password", ""))
    ok = secrets.compare_digest(username, PANEL_USERNAME) and secrets.compare_digest(password, PANEL_PASSWORD)
    if ok:
        request.session["auth"] = True
        return RedirectResponse("/", status_code=302)
    return RedirectResponse("/login?error=1", status_code=302)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=302)


# ---------------- pages ----------------

@app.get("/")
def index(request: Request):
    if not request.session.get("auth"):
        return RedirectResponse("/login")
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


# ---------------- JSON API ----------------

@app.get("/api/stats", dependencies=[Depends(require_auth)])
def api_stats():
    return query_stats()


@app.get("/api/clients", dependencies=[Depends(require_auth)])
def api_clients(q: str = "", date_from: str = "", date_to: str = "",
                 page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=500)):
    return query_clients(q, date_from, date_to, page, page_size)


@app.get("/api/clients/{ip}/info", dependencies=[Depends(require_auth)])
def api_client_info(ip: str):
    return query_client_info(ip)


@app.get("/api/clients/{ip}/domains", dependencies=[Depends(require_auth)])
def api_client_domains(ip: str, search: str = "", date_from: str = "", date_to: str = "",
                         page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=500)):
    return query_client_domains(ip, search, date_from, date_to, page, page_size)


@app.get("/api/search", dependencies=[Depends(require_auth)])
def api_search(q: str, date_from: str = "", date_to: str = "",
                page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=500)):
    return query_search(q, date_from, date_to, page, page_size)


@app.get("/api/recent", dependencies=[Depends(require_auth)])
def api_recent(date_from: str = "", date_to: str = "",
                page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=500)):
    return query_recent(date_from, date_to, page, page_size)


@app.get("/api/contacts", dependencies=[Depends(require_auth)])
def api_list_contacts():
    conn = get_conn()
    rows = [dict(r) for r in conn.execute("SELECT ip, name, source FROM contacts ORDER BY name").fetchall()]
    conn.close()
    for r in rows:
        r["ip_class"] = ip_class(r["ip"])
    return rows


@app.get("/api/vpn-sync/status", dependencies=[Depends(require_auth)])
def api_vpn_sync_status():
    return {
        "enabled": bool(VPN_PANEL_TYPE),
        "panel_type": VPN_PANEL_TYPE or None,
        "interval_seconds": VPN_SYNC_INTERVAL_SECONDS,
    }


@app.post("/api/contacts", dependencies=[Depends(require_auth)])
async def api_upsert_contact(request: Request):
    data = await request.json()
    ip = str(data.get("ip", "")).strip()
    name = str(data.get("name", "")).strip()
    if not ip or not name:
        raise HTTPException(400, "ip and name are required")
    conn = get_conn()
    # a manual add/edit from the panel always wins over VPN sync from now on
    conn.execute(
        """INSERT INTO contacts (ip, name, source) VALUES (?, ?, 'manual')
           ON CONFLICT(ip) DO UPDATE SET name=excluded.name, source='manual'""",
        (ip, name),
    )
    conn.commit()
    conn.close()
    return {"ok": True}


@app.delete("/api/contacts/{ip}", dependencies=[Depends(require_auth)])
def api_delete_contact(ip: str):
    conn = get_conn()
    conn.execute("DELETE FROM contacts WHERE ip=?", (ip,))
    conn.commit()
    conn.close()
    return {"ok": True}


# ---------------- export ----------------

@app.get("/api/export", dependencies=[Depends(require_auth)])
def api_export(
    dataset: str,
    format: str = "csv",
    ip: str = "",
    q: str = "",
    search: str = "",
    date_from: str = "",
    date_to: str = "",
    limit: int = Query(20000, le=100000),
):
    if dataset == "clients":
        columns = ["display_name", "client_ip", "ip_class", "query_count", "domain_count", "first_seen", "last_seen"]
        rows = query_clients(q, date_from, date_to, page_size=limit, paginate=False)["rows"]
        title = "گزارش کاربران"
    elif dataset == "domains":
        if not ip:
            raise HTTPException(400, "ip is required")
        columns = ["qname", "count", "rcode_name", "qtypes", "first_seen", "last_seen"]
        rows = query_client_domains(ip, search, date_from, date_to, page_size=limit, paginate=False)["rows"]
        title = f"دامنه‌های بازدید شده - {ip}"
    elif dataset == "search":
        columns = ["ts", "display_name", "client_ip", "qname", "qtype_name", "rcode_name", "elapsed_ms"]
        rows = query_search(q, date_from, date_to, page_size=limit, paginate=False)["rows"]
        title = f"جستجوی دامنه: {q}"
    elif dataset == "recent":
        columns = ["ts", "display_name", "client_ip", "qname", "qtype_name", "rcode_name", "elapsed_ms"]
        rows = query_recent(date_from, date_to, page_size=limit, paginate=False)["rows"]
        title = "لاگ کوئری‌ها"
    else:
        raise HTTPException(400, "invalid dataset")

    if date_from or date_to:
        title += f"  ({jalali_label(date_from)} تا {jalali_label(date_to)})"

    filename_base = f"{dataset}_{tehran_now().strftime('%Y%m%d_%H%M%S')}"

    if format == "csv":
        return Response(
            rows_to_csv(rows, columns).encode("utf-8"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'},
        )
    if format == "txt":
        return Response(
            rows_to_txt(rows, columns).encode("utf-8"),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.txt"'},
        )
    if format == "xlsx":
        return Response(
            rows_to_xlsx(rows, columns),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'},
        )
    if format == "pdf":
        return Response(
            rows_to_pdf(rows, columns, title),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.pdf"'},
        )
    raise HTTPException(400, "invalid format")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.environ.get("PORT", "8080")))
